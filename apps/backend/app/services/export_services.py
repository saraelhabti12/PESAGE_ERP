from datetime import datetime, timezone
import io
from typing import Any, List, Optional
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.models.credit_note_model import CreditNote
from app.models.employee_model import Employee
from app.models.expense_model import Expense
from app.models.invoice_model import Invoice, InvoiceType
from app.models.partner_model import BusinessPartner, PartnerType
from app.models.payment_model import Payment
from app.models.product_model import Product
from app.models.purchase_model import PurchaseOrder
from app.models.quote_model import Quote
from app.models.sale_model import Sale
from app.models.stock_model import StockBalance, StockMovement
from app.models.vehicle_model import Vehicle, VehicleMaintenance


class ExportService:
    @staticmethod
    def _create_styled_workbook(
        title: str,
        headers: List[str],
        data_rows: List[List[Any]],
        sheet_name: str = "Données",
    ) -> io.BytesIO:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name

        # Brand / Title Banner
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        title_cell = ws.cell(row=1, column=1, value=f"PESAGE 2000 — {title.upper()}")
        title_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
        title_cell.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 32

        # Metadata Row
        now_str = datetime.now(timezone.utc).strftime("%d/%m/%Y à %H:%M UTC")
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
        meta_cell = ws.cell(row=2, column=1, value=f"Document généré le {now_str} — ERP Pesage 2000")
        meta_cell.font = Font(name="Calibri", size=9, italic=True, color="64748B")
        meta_cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[2].height = 18

        # Blank spacing row
        ws.row_dimensions[3].height = 8

        # Header Row (Row 4)
        header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="medium", color="1E293B"),
        )

        for col_idx, h_text in enumerate(headers, start=1):
            cell = ws.cell(row=4, column=col_idx, value=h_text)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

        ws.row_dimensions[4].height = 26

        # Data Rows
        data_font = Font(name="Calibri", size=10, color="0F172A")
        alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        data_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0"),
        )

        for r_idx, row_values in enumerate(data_rows, start=5):
            is_alt = (r_idx % 2 == 0)
            ws.row_dimensions[r_idx].height = 20

            for c_idx, val in enumerate(row_values, start=1):
                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                cell.font = data_font
                cell.border = data_border
                if is_alt:
                    cell.fill = alt_fill

                # Format numbers and currencies
                if isinstance(val, float) or isinstance(val, int):
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    if "DH" in headers[c_idx - 1] or "Montant" in headers[c_idx - 1] or "Prix" in headers[c_idx - 1] or "Total" in headers[c_idx - 1]:
                        cell.number_format = '#,##0.00 "DH"'
                    elif "Qté" in headers[c_idx - 1] or "Stock" in headers[c_idx - 1] or "Kilométrage" in headers[c_idx - 1]:
                        cell.number_format = "#,##0.00"
                elif isinstance(val, datetime):
                    cell.value = val.strftime("%d/%m/%Y")
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        # Auto-fit column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                # ignore merged title / meta row for width calculation
                if cell.row in (1, 2, 3):
                    continue
                v_str = str(cell.value or "")
                if len(v_str) > max_len:
                    max_len = len(v_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    # ----------------------------------------------------
    # DATA EXTRACTORS
    # ----------------------------------------------------
    @classmethod
    def export_partners(cls, db: Session, is_supplier: bool = False) -> io.BytesIO:
        p_types = [PartnerType.SUPPLIER, PartnerType.CUSTOMER_SUPPLIER] if is_supplier else [PartnerType.CUSTOMER, PartnerType.CUSTOMER_SUPPLIER]
        title = "Liste des Fournisseurs" if is_supplier else "Liste des Clients"

        partners = db.query(BusinessPartner).filter(
            BusinessPartner.partner_type.in_(p_types),
            BusinessPartner.is_active == True,
        ).order_by(BusinessPartner.code.asc()).all()

        headers = ["Code", "Raison Sociale", "Nom Commercial", "ICE", "Identifiant Fiscal", "Téléphone", "Email", "Ville", "Date Création"]
        rows = [
            [
                p.code,
                p.legal_name,
                p.trade_name or "",
                p.ice or "",
                p.tax_id or "",
                p.phone or "",
                p.email or "",
                p.city or "",
                p.created_at,
            ]
            for p in partners
        ]
        return cls._create_styled_workbook(title, headers, rows, sheet_name="Partenaires")

    @classmethod
    def export_products(cls, db: Session) -> io.BytesIO:
        products = db.query(Product).filter(Product.is_active == True).order_by(Product.sku.asc()).all()
        headers = ["SKU / Référence", "Désignation", "Code-barres", "Prix Vente HT", "Prix Achat HT", "Taux TVA %", "Suivi Stock", "Catégorie"]
        rows = [
            [
                p.sku,
                p.designation,
                p.barcode or "",
                float(p.sale_price),
                float(p.purchase_price),
                float(p.tax_rate),
                "OUI" if p.track_stock else "NON",
                p.category.name if p.category else "",
            ]
            for p in products
        ]
        return cls._create_styled_workbook("Catalogue Produits", headers, rows, sheet_name="Produits")

    @classmethod
    def export_quotes(cls, db: Session) -> io.BytesIO:
        quotes = db.query(Quote).order_by(Quote.quote_date.desc()).all()
        headers = ["Référence", "Date", "Client", "Statut", "Total HT", "Total TVA", "Total TTC"]
        rows = [
            [
                q.reference,
                q.quote_date,
                q.customer.legal_name if q.customer else "",
                q.status.value,
                float(q.subtotal_ht),
                float(q.tax_amount),
                float(q.total_ttc),
            ]
            for q in quotes
        ]
        return cls._create_styled_workbook("Devis Commerciaux", headers, rows, sheet_name="Devis")

    @classmethod
    def export_sales(cls, db: Session) -> io.BytesIO:
        sales = db.query(Sale).order_by(Sale.order_date.desc()).all()
        headers = ["Référence", "Date", "Client", "Statut", "Total HT", "Total TVA", "Total TTC"]
        rows = [
            [
                s.reference,
                s.order_date,
                s.customer.legal_name if s.customer else "",
                s.status.value,
                float(s.subtotal_ht),
                float(s.tax_amount),
                float(s.total_ttc),
            ]
            for s in sales
        ]
        return cls._create_styled_workbook("Commandes de Vente", headers, rows, sheet_name="Ventes")

    @classmethod
    def export_purchases(cls, db: Session) -> io.BytesIO:
        purchases = db.query(PurchaseOrder).order_by(PurchaseOrder.order_date.desc()).all()
        headers = ["Référence", "Date", "Fournisseur", "Statut", "Total HT", "Total TVA", "Total TTC"]
        rows = [
            [
                p.reference,
                p.order_date,
                p.supplier.legal_name if p.supplier else "",
                p.status.value,
                float(p.subtotal_ht),
                float(p.tax_amount),
                float(p.total_ttc),
            ]
            for p in purchases
        ]
        return cls._create_styled_workbook("Commandes d'Achat", headers, rows, sheet_name="Achats")

    @classmethod
    def export_invoices(cls, db: Session, invoice_type: Optional[InvoiceType] = None) -> io.BytesIO:
        query = db.query(Invoice)
        if invoice_type:
            query = query.filter(Invoice.invoice_type == invoice_type)
        invoices = query.order_by(Invoice.invoice_date.desc()).all()

        title = "Factures Clients" if invoice_type == InvoiceType.CUSTOMER else ("Factures Fournisseurs" if invoice_type == InvoiceType.SUPPLIER else "Toutes les Factures")
        headers = ["Référence", "Type", "Partenaire", "Date Facture", "Date Échéance", "Statut", "Total TTC", "Montant Payé", "Reste à Payer", "Modalité"]
        rows = [
            [
                inv.reference,
                "Client" if inv.invoice_type == InvoiceType.CUSTOMER else "Fournisseur",
                (inv.customer.legal_name if inv.customer else (inv.supplier.legal_name if inv.supplier else "")),
                inv.invoice_date,
                inv.due_date,
                inv.status.value,
                float(inv.total_ttc),
                float(inv.amount_paid),
                float(inv.amount_due),
                inv.payment_terms or "",
            ]
            for inv in invoices
        ]
        return cls._create_styled_workbook(title, headers, rows, sheet_name="Factures")

    @classmethod
    def export_payments(cls, db: Session) -> io.BytesIO:
        payments = db.query(Payment).order_by(Payment.payment_date.desc()).all()
        headers = ["Référence", "Date", "Facture", "Montant", "Mode Règlement", "Réf Transaction"]
        rows = [
            [
                p.reference,
                p.payment_date,
                p.invoice.reference if p.invoice else "",
                float(p.amount),
                p.payment_method,
                p.transaction_reference or "",
            ]
            for p in payments
        ]
        return cls._create_styled_workbook("Règlements & Paiements", headers, rows, sheet_name="Règlements")

    @classmethod
    def export_credit_notes(cls, db: Session) -> io.BytesIO:
        cns = db.query(CreditNote).order_by(CreditNote.credit_note_date.desc()).all()
        headers = ["Référence Avoir", "Date", "Facture Liée", "Client", "Motif", "Statut", "Total TTC"]
        rows = [
            [
                c.reference,
                c.credit_note_date,
                c.invoice.reference if c.invoice else "",
                c.customer.legal_name if c.customer else "",
                c.reason.value,
                c.status.value,
                float(c.total_ttc),
            ]
            for c in cns
        ]
        return cls._create_styled_workbook("Avoirs Clients", headers, rows, sheet_name="Avoirs")

    @classmethod
    def export_stock(cls, db: Session) -> io.BytesIO:
        balances = db.query(StockBalance).filter(StockBalance.quantity > 0).all()
        headers = ["SKU", "Désignation", "Dépôt", "Quantité Totale", "Quantité Réservée", "Quantité Disponible", "Valeur Stock HT"]
        rows = [
            [
                b.product.sku if b.product else "",
                b.product.designation if b.product else "",
                b.warehouse.name if b.warehouse else "",
                float(b.quantity),
                float(b.reserved_quantity),
                float(b.available_quantity),
                round(float(b.quantity) * float(b.product.purchase_price if b.product else 0), 2),
            ]
            for b in balances
        ]
        return cls._create_styled_workbook("État des Stocks en Dépôt", headers, rows, sheet_name="Stocks")

    @classmethod
    def export_stock_movements(cls, db: Session) -> io.BytesIO:
        movements = db.query(StockMovement).order_by(StockMovement.created_at.desc()).limit(1000).all()
        headers = ["Date", "Type Mouvement", "SKU", "Désignation", "Dépôt", "Quantité", "Motif"]
        rows = [
            [
                m.created_at,
                m.movement_type.value,
                m.product.sku if m.product else "",
                m.product.designation if m.product else "",
                m.warehouse.name if m.warehouse else "",
                float(m.quantity),
                m.reason or "",
            ]
            for m in movements
        ]
        return cls._create_styled_workbook("Historique Mouvements de Stock", headers, rows, sheet_name="Mouvements")

    @classmethod
    def export_expenses(cls, db: Session) -> io.BytesIO:
        expenses = db.query(Expense).order_by(Expense.expense_date.desc()).all()
        headers = ["Référence", "Date", "Catégorie", "Libellé", "Montant", "Mode Paiement", "Bénéficiaire", "Statut"]
        rows = [
            [
                e.reference,
                e.expense_date,
                e.category.name if e.category else "",
                e.label,
                float(e.amount),
                e.payment_method or "",
                e.supplier_payee or "",
                e.status.value,
            ]
            for e in expenses
        ]
        return cls._create_styled_workbook("Journal des Dépenses", headers, rows, sheet_name="Dépenses")

    @classmethod
    def export_vehicles(cls, db: Session) -> io.BytesIO:
        vehicles = db.query(Vehicle).order_by(Vehicle.brand.asc()).all()
        headers = ["Code", "Immatriculation", "Marque", "Modèle", "Type", "Kilométrage", "Carburant", "Statut", "Conducteur"]
        rows = [
            [
                v.vehicle_code,
                v.registration_number,
                v.brand,
                v.model,
                v.vehicle_type or "",
                float(v.current_mileage),
                v.fuel_type.value,
                v.status.value,
                f"{v.default_driver.first_name} {v.default_driver.last_name}" if v.default_driver else "",
            ]
            for v in vehicles
        ]
        return cls._create_styled_workbook("Parc Véhicules", headers, rows, sheet_name="Véhicules")

    @classmethod
    def export_maintenances(cls, db: Session) -> io.BytesIO:
        maints = db.query(VehicleMaintenance).order_by(VehicleMaintenance.maintenance_date.desc()).all()
        headers = ["Date", "Véhicule", "Type", "Description", "Kilométrage", "Coût MO", "Coût Pièces", "Coût Total", "Prestataire"]
        rows = [
            [
                m.maintenance_date,
                f"{m.vehicle.brand} {m.vehicle.model} ({m.vehicle.registration_number})" if m.vehicle else "",
                m.maintenance_type.value,
                m.description,
                float(m.mileage),
                float(m.labor_cost),
                float(m.parts_cost),
                float(m.total_cost),
                m.supplier or "",
            ]
            for m in maints
        ]
        return cls._create_styled_workbook("Journal Maintenance Véhicules", headers, rows, sheet_name="Maintenances")

    @classmethod
    def export_employees(cls, db: Session) -> io.BytesIO:
        employees = db.query(Employee).order_by(Employee.employee_code.asc()).all()
        headers = ["Matricule", "Prénom", "Nom", "Téléphone", "Email", "Poste", "Département", "Date d'embauche", "Statut"]
        rows = [
            [
                emp.employee_code,
                emp.first_name,
                emp.last_name,
                emp.phone or "",
                emp.email or "",
                emp.job_title or "",
                emp.department or "",
                emp.hire_date,
                "Actif" if emp.is_active else "Inactif",
            ]
            for emp in employees
        ]
        return cls._create_styled_workbook("Personnel & Employés", headers, rows, sheet_name="Employés")

