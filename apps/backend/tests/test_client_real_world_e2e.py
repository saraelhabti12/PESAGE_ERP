import io
from datetime import datetime, timezone
import pytest
from fastapi import status
from fastapi.testclient import TestClient
import openpyxl

from app.core.config import settings


@pytest.fixture
def auth_headers(client: TestClient):
    res = client.post(
        "/api/v1/auth/login",
        json={"email": settings.ADMIN_EMAIL, "password": settings.ADMIN_PASSWORD},
    )
    assert res.status_code == 200, f"Login failed: {res.text}"
    token = res.json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture
def catalog_prereqs(client: TestClient, auth_headers: dict):
    """Create basic category, product type, and unit for product fixtures."""
    ts = int(datetime.now(timezone.utc).timestamp() * 1000)
    cat = client.post("/api/v1/categories", json={"name": f"Pesage {ts}", "code": f"CAT-{ts}"}, headers=auth_headers).json()
    pt = client.post("/api/v1/product-types", json={"name": f"Industriel {ts}", "code": f"IND-{ts}"}, headers=auth_headers).json()
    u = client.post("/api/v1/units", json={"name": f"Unite {ts}", "abbreviation": f"U{ts}"}, headers=auth_headers).json()
    return {"category_id": cat["id"], "product_type_id": pt["id"], "unit_id": u["id"]}


# ==============================================================================
# SCENARIO A — COMPLETE PURCHASE WORKFLOW
# Supplier → Purchase Order → Approve → Purchase Receipt → Stock Increase in Warehouse
# → Supplier Invoice → Payment → Remaining Amount
# ==============================================================================
def test_scenario_a_purchase_workflow(client: TestClient, auth_headers, catalog_prereqs):
    # 1. Create Supplier
    ts = int(datetime.now(timezone.utc).timestamp() * 1000)
    res_sup = client.post(
        "/api/v1/partners",
        json={
            "legal_name": f"Fournisseur Métrologie Maroc {ts}",
            "partner_type": "SUPPLIER",
            "email": f"contact_{ts}@fournisseur-metrologie.ma",
            "phone": "0522112233",
            "tax_id": f"IF{ts}",
        },
        headers=auth_headers,
    )
    assert res_sup.status_code == status.HTTP_201_CREATED
    supplier = res_sup.json()
    assert supplier["code"].startswith("FOU-")
    supplier_id = supplier["id"]

    # 2. Create Warehouse A
    res_wh_a = client.post(
        "/api/v1/warehouses",
        json={"name": f"Dépôt Principal A {ts}", "code": f"DEP-A-{ts}", "type": "MAIN", "is_active": True},
        headers=auth_headers,
    )
    assert res_wh_a.status_code == status.HTTP_201_CREATED
    wh_a_id = res_wh_a.json()["id"]

    # 3. Create Stock-Managed Product
    res_prod = client.post(
        "/api/v1/products",
        json={
            "sku": f"PES-CAP-{ts}",
            "designation": "Capteur de pesage 50T",
            "category_id": catalog_prereqs["category_id"],
            "product_type_id": catalog_prereqs["product_type_id"],
            "unit_id": catalog_prereqs["unit_id"],
            "purchase_price": 2500.0,
            "sale_price": 4000.0,
            "tax_rate": 20.0,
            "track_stock": True,
        },
        headers=auth_headers,
    )
    assert res_prod.status_code == status.HTTP_201_CREATED
    product_id = res_prod.json()["id"]

    # Check initial stock in Warehouse A is 0
    res_stock_init = client.get(f"/api/v1/stock?product_id={product_id}&warehouse_id={wh_a_id}", headers=auth_headers)
    assert res_stock_init.status_code == status.HTTP_200_OK
    init_balances = res_stock_init.json()
    init_qty = init_balances[0]["quantity"] if init_balances else 0.0
    assert init_qty == 0.0

    # 4. Create Purchase Order (10 units @ 2500 HT)
    res_po = client.post(
        "/api/v1/purchases",
        json={
            "supplier_id": supplier_id,
            "notes": "Commande réapprovisionnement capteurs 50T",
            "items": [
                {
                    "product_id": product_id,
                    "quantity_ordered": 10.0,
                    "unit_price_ht": 2500.0,
                    "tax_rate": 20.0,
                }
            ],
        },
        headers=auth_headers,
    )
    assert res_po.status_code == status.HTTP_201_CREATED
    po = res_po.json()
    po_id = po["id"]
    po_item_id = po["items"][0]["id"]
    assert po["subtotal_ht"] == 25000.0
    assert po["tax_amount"] == 5000.0
    assert po["total_ttc"] == 30000.0

    # CRITICAL RULE: Purchase Order creation MUST NOT increase physical stock
    res_stock_after_po = client.get(f"/api/v1/stock?product_id={product_id}&warehouse_id={wh_a_id}", headers=auth_headers)
    balances_after_po = res_stock_after_po.json()
    qty_after_po = balances_after_po[0]["quantity"] if balances_after_po else 0.0
    assert qty_after_po == 0.0, "Physical stock must not increase on PO creation alone"

    # Approve Purchase Order
    res_appr = client.post(f"/api/v1/purchases/{po_id}/approve", headers=auth_headers)
    assert res_appr.status_code == status.HTTP_200_OK

    # 5. Create Purchase Receipt in Warehouse A
    res_rcpt = client.post(
        "/api/v1/purchase-receipts",
        json={
            "reference": f"BL-FOURN-{ts}",
            "purchase_order_id": po_id,
            "warehouse_id": wh_a_id,
            "notes": "10 capteurs reçus conformes",
            "items": [
                {
                    "purchase_order_item_id": po_item_id,
                    "product_id": product_id,
                    "quantity_received": 10.0,
                }
            ],
        },
        headers=auth_headers,
    )
    assert res_rcpt.status_code == status.HTTP_201_CREATED
    rcpt = res_rcpt.json()
    assert rcpt["warehouse_id"] == wh_a_id

    # 6. Verify Stock Increased in Warehouse A
    res_stock_after_rcpt = client.get(f"/api/v1/stock?product_id={product_id}&warehouse_id={wh_a_id}", headers=auth_headers)
    balances_after_rcpt = res_stock_after_rcpt.json()
    assert len(balances_after_rcpt) > 0
    assert balances_after_rcpt[0]["quantity"] == 10.0

    # 7. Create Supplier Invoice from Purchase
    now = datetime.now(timezone.utc)
    res_supp_inv = client.post(
        "/api/v1/invoices",
        json={
            "supplier_id": supplier_id,
            "invoice_type": "SUPPLIER",
            "invoice_date": now.isoformat(),
            "payment_terms": "30 jours",
            "items": [
                {
                    "product_id": product_id,
                    "description": "Capteur de pesage 50T",
                    "quantity": 10.0,
                    "unit_price_ht": 2500.0,
                    "tax_rate": 20.0,
                }
            ],
        },
        headers=auth_headers,
    )
    assert res_supp_inv.status_code == status.HTTP_201_CREATED
    supp_inv = res_supp_inv.json()
    supp_inv_id = supp_inv["id"]
    assert supp_inv["invoice_type"] == "SUPPLIER"
    assert supp_inv["total_ttc"] == 30000.0
    assert supp_inv["amount_paid"] == 0.0
    assert supp_inv["amount_due"] == 30000.0

    # 8. Record Partial Payment (10,000 DH)
    res_pay1 = client.post(
        "/api/v1/payments",
        json={
            "invoice_id": supp_inv_id,
            "amount": 10000.0,
            "payment_method": "VIREMENT",
            "reference": f"VIR-{ts}-1",
            "payment_date": now.isoformat(),
        },
        headers=auth_headers,
    )
    assert res_pay1.status_code == status.HTTP_201_CREATED

    # Verify Supplier Invoice balances
    res_inv_chk = client.get(f"/api/v1/invoices/{supp_inv_id}", headers=auth_headers)
    assert res_inv_chk.status_code == status.HTTP_200_OK
    inv_data = res_inv_chk.json()
    assert inv_data["amount_paid"] == 10000.0
    assert inv_data["amount_due"] == 20000.0
    assert inv_data["status"] == "PARTIALLY_PAID"

    # 9. Record Final Payment (20,000 DH)
    res_pay2 = client.post(
        "/api/v1/payments",
        json={
            "invoice_id": supp_inv_id,
            "amount": 20000.0,
            "payment_method": "CHEQUE",
            "reference": f"CHQ-{ts}-2",
            "payment_date": now.isoformat(),
        },
        headers=auth_headers,
    )
    assert res_pay2.status_code == status.HTTP_201_CREATED

    # Verify Invoice is now FULLY PAID
    res_inv_paid = client.get(f"/api/v1/invoices/{supp_inv_id}", headers=auth_headers)
    assert res_inv_paid.json()["amount_paid"] == 30000.0
    assert res_inv_paid.json()["amount_due"] == 0.0
    assert res_inv_paid.json()["status"] == "PAID"


# ==============================================================================
# SCENARIO B — COMPLETE SALES WORKFLOW
# Customer → Sale → Source Warehouse → Confirmation → Stock Decrease
# → Delivery Note → Customer Invoice → Payment
# ==============================================================================
def test_scenario_b_sales_workflow(client: TestClient, auth_headers, catalog_prereqs):
    # 1. Create Customer
    ts = int(datetime.now(timezone.utc).timestamp() * 1000)
    res_cust = client.post(
        "/api/v1/partners",
        json={
            "legal_name": f"Société Industrielle Agro {ts}",
            "partner_type": "CUSTOMER",
            "email": f"achats_{ts}@agro-industriel.ma",
            "phone": "0522998877",
            "tax_id": f"IF{ts}",
        },
        headers=auth_headers,
    )
    assert res_cust.status_code == status.HTTP_201_CREATED
    customer = res_cust.json()
    assert customer["code"].startswith("CLI-")
    cust_id = customer["id"]

    # 2. Create Warehouse A and B
    res_wh_a = client.post(
        "/api/v1/warehouses",
        json={"name": f"Magasin Casablanca {ts}", "code": f"MAG-CASA-{ts}", "type": "MAIN", "is_active": True},
        headers=auth_headers,
    )
    wh_a_id = res_wh_a.json()["id"]

    res_wh_b = client.post(
        "/api/v1/warehouses",
        json={"name": f"Magasin Tanger {ts}", "code": f"MAG-TNG-{ts}", "type": "WAREHOUSE", "is_active": True},
        headers=auth_headers,
    )
    wh_b_id = res_wh_b.json()["id"]

    # 3. Create Stock-Managed Product
    res_prod = client.post(
        "/api/v1/products",
        json={
            "sku": f"IND-P2000-{ts}",
            "designation": "Indicateur Numérique P2000-X1",
            "category_id": catalog_prereqs["category_id"],
            "product_type_id": catalog_prereqs["product_type_id"],
            "unit_id": catalog_prereqs["unit_id"],
            "purchase_price": 1200.0,
            "sale_price": 2000.0,
            "tax_rate": 20.0,
            "track_stock": True,
        },
        headers=auth_headers,
    )
    prod_id = res_prod.json()["id"]

    # Initial stock entry: 100 units in Warehouse A, 50 units in Warehouse B via adjustments
    client.post(
        "/api/v1/stock/adjustments",
        json={
            "product_id": prod_id,
            "warehouse_id": wh_a_id,
            "new_quantity": 100.0,
            "reason": "Stock initial Casa",
        },
        headers=auth_headers,
    )
    client.post(
        "/api/v1/stock/adjustments",
        json={
            "product_id": prod_id,
            "warehouse_id": wh_b_id,
            "new_quantity": 50.0,
            "reason": "Stock initial Tanger",
        },
        headers=auth_headers,
    )

    # 4. Create Sale with Source Warehouse = Warehouse A (20 units @ 2000 HT)
    res_sale = client.post(
        "/api/v1/sales",
        json={
            "customer_id": cust_id,
            "notes": "Vente balance pont-bascule",
            "items": [
                {
                    "product_id": prod_id,
                    "warehouse_id": wh_a_id,
                    "quantity": 20.0,
                    "unit_price_ht": 2000.0,
                    "tax_rate": 20.0,
                    "discount_percent": 0.0,
                }
            ],
        },
        headers=auth_headers,
    )
    assert res_sale.status_code == status.HTTP_201_CREATED
    sale = res_sale.json()
    sale_id = sale["id"]
    assert sale["subtotal_ht"] == 40000.0
    assert sale["tax_amount"] == 8000.0
    assert sale["total_ttc"] == 48000.0

    # 5. Confirm Sale -> Stock must be deducted strictly from Warehouse A (100 -> 80), Warehouse B remains 50
    res_conf_sale = client.post(f"/api/v1/sales/{sale_id}/confirm", headers=auth_headers)
    assert res_conf_sale.status_code == status.HTTP_200_OK

    res_bal_a = client.get(f"/api/v1/stock?product_id={prod_id}&warehouse_id={wh_a_id}", headers=auth_headers)
    assert res_bal_a.json()[0]["quantity"] == 80.0, "Warehouse A should decrease from 100 to 80"

    res_bal_b = client.get(f"/api/v1/stock?product_id={prod_id}&warehouse_id={wh_b_id}", headers=auth_headers)
    assert res_bal_b.json()[0]["quantity"] == 50.0, "Warehouse B should remain untouched at 50"

    # 6. Create Delivery Note
    res_dn = client.post(
        "/api/v1/delivery-notes",
        json={
            "sale_id": sale_id,
            "customer_id": cust_id,
            "notes": "Livraison site client Casa",
            "items": [
                {
                    "product_id": prod_id,
                    "warehouse_id": wh_a_id,
                    "delivered_quantity": 20.0,
                }
            ],
        },
        headers=auth_headers,
    )
    assert res_dn.status_code == status.HTTP_201_CREATED

    # Verify no double stock deduction occurred upon delivery note creation
    res_bal_check = client.get(f"/api/v1/stock?product_id={prod_id}&warehouse_id={wh_a_id}", headers=auth_headers)
    assert res_bal_check.json()[0]["quantity"] == 80.0, "Delivery note must not double-deduct stock"

    # 7. Create Customer Invoice from Sale
    now = datetime.now(timezone.utc)
    res_cust_inv = client.post(
        "/api/v1/invoices",
        json={
            "customer_id": cust_id,
            "invoice_type": "CUSTOMER",
            "invoice_date": now.isoformat(),
            "payment_terms": "30 jours",
            "items": [
                {
                    "product_id": prod_id,
                    "description": "Indicateur Numérique P2000-X1",
                    "quantity": 20.0,
                    "unit_price_ht": 2000.0,
                    "tax_rate": 20.0,
                }
            ],
        },
        headers=auth_headers,
    )
    assert res_cust_inv.status_code == status.HTTP_201_CREATED
    cust_inv = res_cust_inv.json()
    cust_inv_id = cust_inv["id"]
    assert cust_inv["invoice_type"] == "CUSTOMER"
    assert cust_inv["total_ttc"] == 48000.0
    assert cust_inv["amount_paid"] == 0.0
    assert cust_inv["amount_due"] == 48000.0

    # 8. Record Partial Payment (20,000 DH)
    res_pay = client.post(
        "/api/v1/payments",
        json={
            "invoice_id": cust_inv_id,
            "amount": 20000.0,
            "payment_method": "VIREMENT",
            "reference": f"VIR-AGRO-{ts}",
            "payment_date": now.isoformat(),
        },
        headers=auth_headers,
    )
    assert res_pay.status_code == status.HTTP_201_CREATED

    res_inv_stat = client.get(f"/api/v1/invoices/{cust_inv_id}", headers=auth_headers)
    assert res_inv_stat.json()["amount_paid"] == 20000.0
    assert res_inv_stat.json()["amount_due"] == 28000.0
    assert res_inv_stat.json()["status"] == "PARTIALLY_PAID"


# ==============================================================================
# SCENARIO C — CREDIT NOTE / AVOIR WORKFLOW
# Invoice → Credit Note → return_to_stock → Stock Increase & Intact Original Invoice
# ==============================================================================
def test_scenario_c_credit_note_workflow(client: TestClient, auth_headers, catalog_prereqs):
    ts = int(datetime.now(timezone.utc).timestamp() * 1000)
    res_cust = client.post(
        "/api/v1/partners",
        json={"legal_name": f"Client Avoir Test SARL {ts}", "partner_type": "CUSTOMER"},
        headers=auth_headers,
    )
    cust_id = res_cust.json()["id"]

    res_wh = client.post(
        "/api/v1/warehouses",
        json={"name": f"Entrepôt Retour {ts}", "code": f"DEP-RET-{ts}", "type": "MAIN", "is_active": True},
        headers=auth_headers,
    )
    wh_id = res_wh.json()["id"]

    res_prod = client.post(
        "/api/v1/products",
        json={
            "sku": f"ACC-CABLE-{ts}",
            "designation": "Câble Blindé 10M",
            "category_id": catalog_prereqs["category_id"],
            "product_type_id": catalog_prereqs["product_type_id"],
            "unit_id": catalog_prereqs["unit_id"],
            "purchase_price": 200.0,
            "sale_price": 500.0,
            "tax_rate": 20.0,
            "track_stock": True,
        },
        headers=auth_headers,
    )
    prod_id = res_prod.json()["id"]

    # Initial stock in warehouse is 10 units
    client.post(
        "/api/v1/stock/adjustments",
        json={
            "product_id": prod_id,
            "warehouse_id": wh_id,
            "new_quantity": 10.0,
            "reason": "Stock initial avant vente",
        },
        headers=auth_headers,
    )

    now = datetime.now(timezone.utc)
    res_inv = client.post(
        "/api/v1/invoices",
        json={
            "customer_id": cust_id,
            "invoice_type": "CUSTOMER",
            "invoice_date": now.isoformat(),
            "items": [
                {
                    "product_id": prod_id,
                    "description": "Câble Blindé 10M",
                    "quantity": 10.0,
                    "unit_price_ht": 500.0,
                    "tax_rate": 20.0,
                }
            ],
        },
        headers=auth_headers,
    )
    assert res_inv.status_code == status.HTTP_201_CREATED
    inv_id = res_inv.json()["id"]
    inv_item_id = res_inv.json()["items"][0]["id"]
    assert res_inv.json()["total_ttc"] == 6000.0

    # 2. Create Partial Credit Note (Return 3 units with return_to_stock=True)
    res_cn = client.post(
        "/api/v1/credit-notes",
        json={
            "invoice_id": inv_id,
            "reason": "RETURN",
            "notes": "3 câbles défectueux retournés",
            "items": [
                {
                    "product_id": prod_id,
                    "invoice_item_id": inv_item_id,
                    "warehouse_id": wh_id,
                    "description": "Retour 3 câbles défectueux",
                    "quantity": 3.0,
                    "unit_price_ht": 500.0,
                    "tax_rate": 20.0,
                    "return_to_stock": True,
                }
            ],
        },
        headers=auth_headers,
    )
    assert res_cn.status_code == status.HTTP_201_CREATED
    cn = res_cn.json()
    assert cn["subtotal_ht"] == 1500.0
    assert cn["tax_amount"] == 300.0
    assert cn["total_ttc"] == 1800.0

    # 3. Stock increases by 3 upon return_to_stock (10 + 3 = 13)
    res_stock = client.get(f"/api/v1/stock?product_id={prod_id}&warehouse_id={wh_id}", headers=auth_headers)
    assert res_stock.json()[0]["quantity"] == 13.0, "Stock must increase upon return_to_stock credit note"

    # 4. Verify Original Invoice is INTACT and NOT deleted/destroyed
    res_inv_check = client.get(f"/api/v1/invoices/{inv_id}", headers=auth_headers)
    assert res_inv_check.status_code == status.HTTP_200_OK
    assert res_inv_check.json()["total_ttc"] == 6000.0

    # 5. Overcredit Rule: Creating credit for 8 more units (total 3 + 8 = 11 > 10) must be rejected
    res_cn_over = client.post(
        "/api/v1/credit-notes",
        json={
            "invoice_id": inv_id,
            "reason": "COMMERCIAL_DISCOUNT",
            "notes": "Tentative sur-crédit",
            "items": [
                {
                    "product_id": prod_id,
                    "invoice_item_id": inv_item_id,
                    "warehouse_id": wh_id,
                    "description": "Sur-crédit",
                    "quantity": 8.0,
                    "unit_price_ht": 500.0,
                    "tax_rate": 20.0,
                    "return_to_stock": False,
                }
            ],
        },
        headers=auth_headers,
    )
    assert res_cn_over.status_code in (status.HTTP_400_BAD_REQUEST, status.HTTP_422_UNPROCESSABLE_ENTITY), "Cannot credit more quantity than originally invoiced"


# ==============================================================================
# SCENARIO D — EXPENSES & RECURRING EXPENSES WORKFLOW & STATS
# ==============================================================================
def test_scenario_d_expenses_and_stats(client: TestClient, auth_headers):
    # 1. Zero Expenses stats check -> Must return 200 with 0.0 values, NOT 422 or 500
    res_zero_stats = client.get("/api/v1/expenses/stats", headers=auth_headers)
    assert res_zero_stats.status_code == status.HTTP_200_OK
    data_zero = res_zero_stats.json()
    assert data_zero["total_expenses"] == 0.0
    assert data_zero["fixed_expenses"] == 0.0
    assert data_zero["variable_expenses"] == 0.0

    # 2. Get/Create Expense Category
    res_cats = client.get("/api/v1/expenses/categories", headers=auth_headers)
    assert res_cats.status_code == status.HTTP_200_OK
    cats = res_cats.json()
    if cats:
        cat_id = cats[0]["id"]
    else:
        res_c = client.post("/api/v1/expenses/categories", json={"name": "Locaux & Bureaux", "code": "LOC-01"}, headers=auth_headers)
        cat_id = res_c.json()["id"]

    # 3. Create Fixed Monthly Expense
    res_exp1 = client.post(
        "/api/v1/expenses",
        json={
            "category_id": cat_id,
            "label": "Loyer Entrepôt Juillet 2025",
            "amount": 15000.0,
            "payment_method": "Virement bancaire",
            "expense_date": "2025-07-05T10:00:00Z",
        },
        headers=auth_headers,
    )
    assert res_exp1.status_code == status.HTTP_201_CREATED

    # 4. Create Variable Expense in August 2026
    res_exp2 = client.post(
        "/api/v1/expenses",
        json={
            "category_id": cat_id,
            "label": "Achat petit outillage Août 2026",
            "amount": 3500.0,
            "payment_method": "Carte bancaire",
            "expense_date": "2026-08-15T14:30:00Z",
        },
        headers=auth_headers,
    )
    assert res_exp2.status_code == status.HTTP_201_CREATED

    # 5. Filter Expenses by Period boundaries
    res_2025 = client.get("/api/v1/expenses?date_from=2025-01-01T00:00:00Z&date_to=2025-12-31T23:59:59Z", headers=auth_headers)
    assert res_2025.status_code == status.HTTP_200_OK
    items_2025 = res_2025.json()
    assert len(items_2025) >= 1
    assert any("2025" in item["label"] for item in items_2025)
    assert not any("2026" in item["label"] for item in items_2025)

    res_aug2026 = client.get("/api/v1/expenses?date_from=2026-08-01T00:00:00Z&date_to=2026-08-31T23:59:59Z", headers=auth_headers)
    assert res_aug2026.status_code == status.HTTP_200_OK
    items_aug2026 = res_aug2026.json()
    assert len(items_aug2026) >= 1
    assert any("2026" in item["label"] for item in items_aug2026)
    assert not any("2025" in item["label"] for item in items_aug2026)


# ==============================================================================
# SCENARIO E — MULTI-WAREHOUSE STOCK MOVEMENTS & TRANSFERS
# ==============================================================================
def test_scenario_e_multi_warehouse_transfers(client: TestClient, auth_headers, catalog_prereqs):
    ts = int(datetime.now(timezone.utc).timestamp() * 1000)
    # 1. Create Warehouse A and B
    res_wh_a = client.post(
        "/api/v1/warehouses",
        json={"name": f"Site A - Casablanca {ts}", "code": f"SITE-A-{ts}", "type": "MAIN", "is_active": True},
        headers=auth_headers,
    )
    wh_a_id = res_wh_a.json()["id"]

    res_wh_b = client.post(
        "/api/v1/warehouses",
        json={"name": f"Site B - Marrakech {ts}", "code": f"SITE-B-{ts}", "type": "WAREHOUSE", "is_active": True},
        headers=auth_headers,
    )
    wh_b_id = res_wh_b.json()["id"]

    # 2. Create Product X
    res_prod = client.post(
        "/api/v1/products",
        json={
            "sku": f"MOD-GSM-{ts}",
            "designation": "Module GSM 4G Pesage",
            "category_id": catalog_prereqs["category_id"],
            "product_type_id": catalog_prereqs["product_type_id"],
            "unit_id": catalog_prereqs["unit_id"],
            "purchase_price": 800.0,
            "sale_price": 1400.0,
            "tax_rate": 20.0,
            "track_stock": True,
        },
        headers=auth_headers,
    )
    prod_id = res_prod.json()["id"]

    # 3. Initial stock: A = 100, B = 50
    client.post(
        "/api/v1/stock/adjustments",
        json={
            "product_id": prod_id,
            "warehouse_id": wh_a_id,
            "new_quantity": 100.0,
            "reason": "Stock initial Site A",
        },
        headers=auth_headers,
    )
    client.post(
        "/api/v1/stock/adjustments",
        json={
            "product_id": prod_id,
            "warehouse_id": wh_b_id,
            "new_quantity": 50.0,
            "reason": "Stock initial Site B",
        },
        headers=auth_headers,
    )

    # 4. Transfer 30 from A to B
    res_transfer = client.post(
        "/api/v1/stock/transfers",
        json={
            "product_id": prod_id,
            "source_warehouse_id": wh_a_id,
            "dest_warehouse_id": wh_b_id,
            "quantity": 30.0,
            "reason": "Transfert Casa vers Marrakech",
        },
        headers=auth_headers,
    )
    assert res_transfer.status_code == status.HTTP_200_OK

    # Expected: A = 70, B = 80
    res_bal_a = client.get(f"/api/v1/stock?product_id={prod_id}&warehouse_id={wh_a_id}", headers=auth_headers)
    assert res_bal_a.json()[0]["quantity"] == 70.0

    res_bal_b = client.get(f"/api/v1/stock?product_id={prod_id}&warehouse_id={wh_b_id}", headers=auth_headers)
    assert res_bal_b.json()[0]["quantity"] == 80.0

    # 5. Validation: Transfer to same warehouse must fail
    res_same = client.post(
        "/api/v1/stock/transfers",
        json={
            "product_id": prod_id,
            "source_warehouse_id": wh_a_id,
            "dest_warehouse_id": wh_a_id,
            "quantity": 10.0,
        },
        headers=auth_headers,
    )
    assert res_same.status_code in (status.HTTP_400_BAD_REQUEST, status.HTTP_422_UNPROCESSABLE_ENTITY)

    # 6. Validation: Transfer more than available must fail
    res_excess = client.post(
        "/api/v1/stock/transfers",
        json={
            "product_id": prod_id,
            "source_warehouse_id": wh_a_id,
            "dest_warehouse_id": wh_b_id,
            "quantity": 150.0,
        },
        headers=auth_headers,
    )
    assert res_excess.status_code in (status.HTTP_400_BAD_REQUEST, status.HTTP_422_UNPROCESSABLE_ENTITY)

    # 7. Validation: Zero or negative quantity must fail
    res_neg = client.post(
        "/api/v1/stock/transfers",
        json={
            "product_id": prod_id,
            "source_warehouse_id": wh_a_id,
            "dest_warehouse_id": wh_b_id,
            "quantity": -5.0,
        },
        headers=auth_headers,
    )
    assert res_neg.status_code in (status.HTTP_400_BAD_REQUEST, status.HTTP_422_UNPROCESSABLE_ENTITY)


# ==============================================================================
# STRICT ISOLATION: CUSTOMER INVOICES VS SUPPLIER INVOICES
# ==============================================================================
def test_invoice_type_strict_backend_isolation(client: TestClient, auth_headers):
    ts = int(datetime.now(timezone.utc).timestamp() * 1000)
    # Create Customer
    res_c = client.post("/api/v1/partners", json={"legal_name": f"Client Isolation SARL {ts}", "partner_type": "CUSTOMER"}, headers=auth_headers)
    cust_id = res_c.json()["id"]

    # Create Supplier
    res_s = client.post("/api/v1/partners", json={"legal_name": f"Fournisseur Isolation SARL {ts}", "partner_type": "SUPPLIER"}, headers=auth_headers)
    supp_id = res_s.json()["id"]

    # Create Customer Invoice
    client.post(
        "/api/v1/invoices",
        json={
            "customer_id": cust_id,
            "invoice_type": "CUSTOMER",
            "items": [{"description": "Service A", "quantity": 1, "unit_price_ht": 1000, "tax_rate": 20}],
        },
        headers=auth_headers,
    )

    # Create Supplier Invoice
    client.post(
        "/api/v1/invoices",
        json={
            "supplier_id": supp_id,
            "invoice_type": "SUPPLIER",
            "items": [{"description": "Achat B", "quantity": 1, "unit_price_ht": 2000, "tax_rate": 20}],
        },
        headers=auth_headers,
    )

    # Fetch with invoice_type=CUSTOMER -> Only customer invoices
    res_cust_list = client.get("/api/v1/invoices?invoice_type=CUSTOMER", headers=auth_headers)
    assert res_cust_list.status_code == status.HTTP_200_OK
    cust_invoices = res_cust_list.json()
    assert len(cust_invoices) >= 1
    assert all(inv["invoice_type"] == "CUSTOMER" for inv in cust_invoices)

    # Fetch with invoice_type=SUPPLIER -> Only supplier invoices
    res_supp_list = client.get("/api/v1/invoices?invoice_type=SUPPLIER", headers=auth_headers)
    assert res_supp_list.status_code == status.HTTP_200_OK
    supp_invoices = res_supp_list.json()
    assert len(supp_invoices) >= 1
    assert all(inv["invoice_type"] == "SUPPLIER" for inv in supp_invoices)


# ==============================================================================
# EXCEL EXPORT VALIDATION & SECURITY (NO SECRETS EXPOSED)
# ==============================================================================
def test_excel_exports_integrity_and_security(client: TestClient, auth_headers):
    export_endpoints = [
        "/api/v1/exports/customers",
        "/api/v1/exports/suppliers",
        "/api/v1/exports/employees",
        "/api/v1/exports/products",
        "/api/v1/exports/quotes",
        "/api/v1/exports/sales",
        "/api/v1/exports/purchases",
        "/api/v1/exports/invoices",
        "/api/v1/exports/invoices/customers",
        "/api/v1/exports/invoices/suppliers",
        "/api/v1/exports/credit-notes",
        "/api/v1/exports/payments",
        "/api/v1/exports/stock",
        "/api/v1/exports/stock/movements",
        "/api/v1/exports/expenses",
        "/api/v1/exports/vehicles",
        "/api/v1/exports/vehicles/maintenances",
    ]

    for ep in export_endpoints:
        res = client.get(ep, headers=auth_headers)
        assert res.status_code == status.HTTP_200_OK, f"Export endpoint {ep} failed: {res.text}"
        assert res.headers.get("content-type") == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        
        # Verify openpyxl can load the generated Excel file cleanly
        excel_bytes = io.BytesIO(res.content)
        wb = openpyxl.load_workbook(excel_bytes)
        assert len(wb.sheetnames) > 0
        sheet = wb.active
        
        # Security sanitization: Scan all cell values to ensure no passwords/hashes/tokens exist
        for row in sheet.iter_rows(values_only=True):
            for cell in row:
                if isinstance(cell, str):
                    assert not cell.startswith("$2b$"), "Password hash leaked into Excel export!"
                    assert not cell.startswith("$argon2"), "Password hash leaked into Excel export!"
                    assert not cell.startswith("eyJ"), "JWT token leaked into Excel export!"
                    assert "SECRET_KEY" not in cell, "Secret key leaked into Excel export!"


# ==============================================================================
# DASHBOARD REAL STATISTICS INTEGRITY
# ==============================================================================
def test_dashboard_stats_real_data(client: TestClient, auth_headers, catalog_prereqs):
    ts = int(datetime.now(timezone.utc).timestamp() * 1000)
    # 1. Zero state
    res_zero = client.get("/api/v1/dashboard/stats", headers=auth_headers)
    assert res_zero.status_code == status.HTTP_200_OK
    d_zero = res_zero.json()
    assert "partners" in d_zero
    assert "purchases" in d_zero
    assert "sales" in d_zero

    # 2. Create customer, warehouse, and sale
    res_c = client.post("/api/v1/partners", json={"legal_name": f"Client Dashboard SARL {ts}", "partner_type": "CUSTOMER"}, headers=auth_headers)
    cust_id = res_c.json()["id"]

    res_wh = client.post("/api/v1/warehouses", json={"name": f"Depot Dash {ts}", "code": f"DEP-DASH-{ts}", "type": "MAIN"}, headers=auth_headers)
    wh_id = res_wh.json()["id"]

    res_p = client.post(
        "/api/v1/products",
        json={
            "sku": f"DASH-{ts}",
            "designation": "Prod Dash",
            "category_id": catalog_prereqs["category_id"],
            "product_type_id": catalog_prereqs["product_type_id"],
            "unit_id": catalog_prereqs["unit_id"],
            "purchase_price": 100,
            "sale_price": 500,
            "tax_rate": 20,
            "track_stock": True,
        },
        headers=auth_headers,
    )
    prod_id = res_p.json()["id"]

    # Initial stock entry
    client.post(
        "/api/v1/stock/adjustments",
        json={
            "product_id": prod_id,
            "warehouse_id": wh_id,
            "new_quantity": 100.0,
            "reason": "Stock initial dashboard",
        },
        headers=auth_headers,
    )

    # Make sale of 5000 DH HT (6000 TTC)
    res_sale = client.post(
        "/api/v1/sales",
        json={
            "customer_id": cust_id,
            "items": [{"product_id": prod_id, "warehouse_id": wh_id, "quantity": 10, "unit_price_ht": 500, "tax_rate": 20}],
        },
        headers=auth_headers,
    )
    assert res_sale.status_code == status.HTTP_201_CREATED
    sale_id = res_sale.json()["id"]
    res_conf = client.post(f"/api/v1/sales/{sale_id}/confirm", headers=auth_headers)
    assert res_conf.status_code == status.HTTP_200_OK

    # 3. Check dashboard reflects real data
    res_after = client.get("/api/v1/dashboard/stats", headers=auth_headers)
    assert res_after.status_code == status.HTTP_200_OK
    d_after = res_after.json()
    assert d_after["partners"]["total_customers"] >= 1
    assert d_after["sales"]["total_ttc"] >= 6000.0
